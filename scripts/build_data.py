#!/usr/bin/env python3
"""Build the data files for the Muriel Lake water level page.

Pulls water levels for Water Survey of Canada station 06AC007 (Muriel Lake
near Gurneyville, AB), merges the approved historical record with the
provisional real-time record, and writes:

  site/data/lake.json        the merged daily series (the page does all the maths)
  site/data/stations.json    every hydrometric station, for the page's station search
  site/data/muriel-lake-water-levels.xlsx
                             full tables, projections and charts for download

Data sources
  * Approved historical daily means (HYDAT), via the GeoMet OGC API:
      https://api.weather.gc.ca/collections/hydrometric-daily-mean
  * Provisional real-time readings (about the last 18 months), via WaterOffice:
      https://wateroffice.ec.gc.ca/services/real_time_data/csv/inline

Approved data always wins. Real-time data only fills dates the approved
record doesn't have yet, so when Environment Canada publishes a new year of
approved data the page switches to it automatically on the next run.

Needs openpyxl for the Excel file; everything else is the standard library.
"""

import bisect
import csv
import datetime as dt
import io
import json
import statistics
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

STATION = "06AC007"
STATION_NAME = "Muriel Lake near Gurneyville"
HYDAT_URL = "https://api.weather.gc.ca/collections/hydrometric-daily-mean/items"
STATIONS_URL = "https://api.weather.gc.ca/collections/hydrometric-stations/items"
REALTIME_URL = "https://wateroffice.ec.gc.ca/services/real_time_data/csv/inline"
PAGE_SIZE = 10000
REALTIME_MONTHS = 19          # WaterOffice keeps roughly 18 months
LST_OFFSET = dt.timedelta(hours=7)   # Alberta local standard time = UTC-7 (what HYDAT uses)
QUARTER_MONTHS = (1, 4, 7, 10)
QUARTER_NAMES = ("Jan 1", "Apr 1", "Jul 1", "Oct 1")
MONTH_NAMES = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
PROJECTION_YEARS = 50
OUT = Path(__file__).resolve().parent.parent / "site" / "data" / "lake.json"
XLSX = OUT.with_name("muriel-lake-water-levels.xlsx")
STATIONS_OUT = OUT.with_name("stations.json")
LAT, LON = 54.1446, -110.7440
IN_PER_M = 100 / 2.54


def fetch(url, params, tries=4):
    full = url + "?" + urllib.parse.urlencode(params, safe="[]:,", quote_via=urllib.parse.quote)
    req = urllib.request.Request(full, headers={"User-Agent": "muriel-lake-levels (github.com/jaybeehan/muriel-lake)"})
    for attempt in range(tries):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.read().decode("utf-8-sig", errors="replace")
        except (urllib.error.URLError, TimeoutError) as e:
            if attempt == tries - 1:
                raise
            print(f"  retrying after error: {e}", file=sys.stderr)
            time.sleep(10 * (attempt + 1))


# ---------------------------------------------------------------- fetching

def fetch_hydat():
    """Approved daily means: {date: level_m}."""
    out = {}
    offset = 0
    while True:
        text = fetch(HYDAT_URL, {
            "STATION_NUMBER": STATION, "f": "csv", "limit": PAGE_SIZE,
            "offset": offset, "sortby": "DATE", "properties": "DATE,LEVEL",
        })
        rows = list(csv.DictReader(io.StringIO(text)))
        for row in rows:
            d, lev = (row.get("DATE") or "").strip(), (row.get("LEVEL") or "").strip()
            if d and lev:
                try:
                    out[dt.date.fromisoformat(d[:10])] = float(lev)
                except ValueError:
                    pass
        print(f"  HYDAT offset {offset}: {len(rows)} rows")
        if len(rows) < PAGE_SIZE:
            break
        offset += PAGE_SIZE
    return out


def fetch_realtime(today):
    """Provisional readings averaged to local-standard-time days: {date: level_m}."""
    sums, counts = {}, {}
    first_of_month = today.replace(day=1)
    for k in range(REALTIME_MONTHS):
        y, m = first_of_month.year, first_of_month.month - k
        while m < 1:
            m += 12
            y -= 1
        start = dt.date(y, m, 1)
        end = dt.date(y + (m == 12), m % 12 + 1, 1)
        try:
            text = fetch(REALTIME_URL, {
                "stations[]": STATION, "parameters[]": 46,
                "start_date": f"{start.isoformat()} 07:00:00",
                "end_date": f"{end.isoformat()} 07:00:00",
            })
        except Exception as e:  # a missing month shouldn't sink the whole run
            print(f"  real-time {start:%Y-%m}: failed ({e})", file=sys.stderr)
            continue
        reader = csv.reader(io.StringIO(text))
        next(reader, None)  # header
        n = 0
        for row in reader:
            if len(row) < 4:
                continue
            try:
                ts = dt.datetime.fromisoformat(row[1].strip().replace("Z", "+00:00"))
                val = float(row[3])
            except ValueError:
                continue
            day = (ts.replace(tzinfo=None) - LST_OFFSET).date()
            if day >= end:   # the closing reading belongs to next month's chunk
                continue
            sums[day] = sums.get(day, 0.0) + val
            counts[day] = counts.get(day, 0) + 1
            n += 1
        print(f"  real-time {start:%Y-%m}: {n} readings")
    return {d: sums[d] / counts[d] for d in sums}


def fetch_stations():
    """All hydrometric stations: [number, name, province, lat, lon, active, real_time]."""
    text = fetch(STATIONS_URL, {"f": "csv", "limit": 20000,
                                "properties": "STATION_NUMBER,STATION_NAME,PROV_TERR_STATE_LOC,STATUS_EN,REAL_TIME"})
    out = []
    for row in csv.DictReader(io.StringIO(text)):
        try:
            out.append([row["STATION_NUMBER"], row["STATION_NAME"].strip(), row["PROV_TERR_STATE_LOC"],
                        round(float(row["y"]), 4), round(float(row["x"]), 4),
                        1 if row.get("STATUS_EN") == "Active" else 0, 1 if row.get("REAL_TIME") == "1" else 0])
        except (KeyError, ValueError):
            continue
    out.sort(key=lambda r: r[0])
    return out


# ---------------------------------------------------------------- maths

def merge(hydat, realtime):
    """Approved data first; real-time only for dates HYDAT doesn't have."""
    series = {d: (v, "approved") for d, v in hydat.items()}
    for d, v in realtime.items():
        if d not in series:
            series[d] = (v, "provisional")
    dates = sorted(series)
    return dates, [series[d][0] for d in dates], [series[d][1] for d in dates]


def interpolate(dates, levels, sources, target):
    """Level on `target`: measured if present, otherwise straight line between neighbours."""
    i = bisect.bisect_left(dates, target)
    if i < len(dates) and dates[i] == target:
        return levels[i], "measured" if sources[i] == "approved" else "measured (provisional)"
    if i == 0 or i == len(dates):
        return None, None
    d0, d1 = dates[i - 1], dates[i]
    v0, v1 = levels[i - 1], levels[i]
    frac = (target - d0).days / (d1 - d0).days
    return v0 + (v1 - v0) * frac, "interpolated"


def percentile(values, p):
    """Same as Google Sheets PERCENTILE (inclusive, linear)."""
    s = sorted(values)
    pos = (len(s) - 1) * p
    lo = int(pos)
    hi = min(lo + 1, len(s) - 1)
    return s[lo] + (s[hi] - s[lo]) * (pos - lo)


def linear_fit(xs, ys):
    mx, my = statistics.fmean(xs), statistics.fmean(ys)
    sxx = sum((x - mx) ** 2 for x in xs)
    slope = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sxx
    return slope, my - slope * mx


def r(x, n=3):
    return None if x is None else round(x, n)


def build():
    today = (dt.datetime.now(dt.timezone.utc).replace(tzinfo=None) - LST_OFFSET).date()
    print("Fetching approved historical data…")
    hydat = fetch_hydat()
    if not hydat:
        sys.exit("No historical data came back; leaving the old data in place.")
    print("Fetching real-time data…")
    realtime = fetch_realtime(today)

    dates, levels, sources = merge(hydat, realtime)
    print(f"Merged series: {len(dates)} days, {dates[0]} to {dates[-1]}")

    # ---- level on given days each year, plus total change (since the first value in each column)
    #      and change vs the previous year, in cm
    def table(months):
        rows = []
        for year in range(dates[0].year, dates[-1].year + 1):
            vals, srcs = [], []
            for m in months:
                v, s = interpolate(dates, levels, sources, dt.date(year, m, 1))
                vals.append(r(v))
                srcs.append(s)
            rows.append({"year": year, "level": vals, "source": srcs})
        n = len(months)
        first = [next((q["level"][c] for q in rows if q["level"][c] is not None), None) for c in range(n)]
        for i, q in enumerate(rows):
            q["total_cm"] = [r((v - first[c]) * 100, 1) if v is not None else None for c, v in enumerate(q["level"])]
            prev = rows[i - 1]["level"] if i else [None] * n
            q["yoy_cm"] = [r((v - p) * 100, 1) if v is not None and p is not None else None
                           for v, p in zip(q["level"], prev)]
        return rows

    quarterly = table(QUARTER_MONTHS)
    monthly = table(range(1, 13))

    # ---- July 1 projection (same method as the Google Sheet)
    last = quarterly[-1]
    anchor_c = max(c for c in range(4) if last["level"][c] is not None)
    anchor_level = last["level"][anchor_c]
    shift = 1 if anchor_c >= 2 else 0          # anchored on Jul/Oct -> project next year's July
    target_year = last["year"] + shift
    deltas = []
    for i, q in enumerate(quarterly):
        j = i + shift
        if j < len(quarterly):
            a, b = q["level"][anchor_c], quarterly[j]["level"][2]
            if a is not None and b is not None:
                deltas.append(b - a)
    jul = [(q["year"], q["level"][2]) for q in quarterly if q["level"][2] is not None]
    last_jul_year, last_jul = jul[-1]
    scen = []
    for name, p in [("Driest on record", 0), ("Dry year", .10), ("Below normal", .25), ("Typical (median)", .50),
                    ("Above normal", .75), ("Wet year", .90), ("Wettest on record", 1)]:
        lev = anchor_level + percentile(deltas, p)
        scen.append({"name": name, "percentile": round(p * 100), "level": r(lev),
                     "change_cm": r((lev - anchor_level) * 100, 1),
                     "vs_last_jul_cm": r((lev - last_jul) * 100, 1)})

    ys, vs = [y for y, _ in jul], [v for _, v in jul]
    s_all, i_all = linear_fit(ys, vs)
    s_10, i_10 = linear_fit(ys[-10:], vs[-10:])
    med_change = statistics.median(b - a for a, b in zip(vs, vs[1:]))
    long_term = []
    for y in range(target_year, target_year + PROJECTION_YEARS):
        long_term.append({"year": y,
                          "trend_all": r(s_all * y + i_all),
                          "trend_10": r(s_10 * y + i_10),
                          "typical": r(last_jul + med_change * (y - last_jul_year))})

    # ---- level relative to earliest quarterly reading
    qbase = next((q["year"], c, q["level"][c]) for q in quarterly for c in range(4) if q["level"][c] is not None)

    data = {
        "station": STATION,
        "station_name": STATION_NAME,
        "updated": dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None).isoformat() + "Z",
        "latest": {"date": dates[-1].isoformat(), "level": r(levels[-1]),
                   "source": "approved" if sources[-1] == "approved" else "provisional"},
        "approved_until": max(hydat).isoformat(),
        "quarter_names": QUARTER_NAMES,
        "quarterly": quarterly,
        "month_names": MONTH_NAMES,
        "monthly": monthly,
        "quarterly_base": {"date": dt.date(qbase[0], QUARTER_MONTHS[qbase[1]], 1).isoformat(), "level": qbase[2]},
        "daily": {
            "base_date": dates[0].isoformat(),
            "base_level": r(levels[0]),
            "dates": [d.isoformat() for d in dates],
            "rel_m": [r(v - levels[0]) for v in levels],
        },
        "projection": {
            "anchor": {"date": dt.date(last["year"], QUARTER_MONTHS[anchor_c], 1).isoformat(), "level": anchor_level},
            "target": dt.date(target_year, 7, 1).isoformat(),
            "years_of_history": len(deltas),
            "last_july": {"year": last_jul_year, "level": last_jul},
            "scenarios": scen,
            "median_july_change_cm": r(med_change * 100, 1),
            "long_term": long_term,
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    # compact daily series for the page: day offsets from t0, levels, provisional flags
    t0 = dates[0]
    page = {
        "station": STATION, "station_name": STATION_NAME, "lat": LAT, "lon": LON,
        "updated": data["updated"], "approved_until": data["approved_until"],
        "realtime_note": "Provisional real-time data covers about the last 18 months.",
        "t0": t0.isoformat(),
        "d": [(d - t0).days for d in dates],
        "v": [r(v) for v in levels],
        "p": "".join("0" if s_ == "approved" else "1" for s_ in sources),
    }
    OUT.write_text(json.dumps(page, separators=(",", ":")))
    print(f"Wrote {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")
    try:
        st = fetch_stations()
        if len(st) > 1000:
            STATIONS_OUT.write_text(json.dumps(st, separators=(",", ":"), ensure_ascii=False))
            print(f"Wrote {STATIONS_OUT} ({len(st)} stations, {STATIONS_OUT.stat().st_size / 1024:.0f} KB)")
        else:
            print(f"Station list looked incomplete ({len(st)} rows); skipped", file=sys.stderr)
    except Exception as e:  # the page falls back to searching the API directly
        print(f"Station list failed: {e}", file=sys.stderr)
    write_xlsx(data, dates, levels, sources)


# ---------------------------------------------------------------- Excel download

def write_xlsx(data, dates, levels, sources):
    """The full tables as an Excel workbook (site/data/muriel-lake-water-levels.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.chart.marker import DataPoint

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE8F5")
    rise_fill = PatternFill("solid", fgColor="DCF1E2")
    rise_font = Font(bold=True, color="1E7B3C")
    interp_font = Font(italic=True, color="8A94A3")
    wb = Workbook()

    def header(ws, cols, widths=None):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 12)
        ws.freeze_panes = "B2"

    # Daily
    ws = wb.active
    ws.title = "Daily"
    header(ws, ["Date", "Level (m above sea level)", "vs first reading (m)", "vs first reading (ft)",
                "vs first reading (in)", "Source"], {1: 12, 2: 16, 3: 14, 4: 14, 5: 14, 6: 26})
    base = levels[0]
    for d, v, s in zip(dates, levels, sources):
        ws.append([d, round(v, 3), round(v - base, 3), round((v - base) / 0.3048, 3), round((v - base) * IN_PER_M, 1),
                   "Approved (HYDAT)" if s == "approved" else "Provisional (real-time)"])
    for row in ws.iter_rows(min_row=2, max_col=5):
        row[0].number_format = "yyyy-mm-dd"
        row[1].number_format = row[2].number_format = row[3].number_format = "0.000"
        row[4].number_format = "0.0"
    ws_daily, n_daily = ws, len(dates)

    months = [f"{m} 1" for m in MONTH_NAMES]
    M = data["monthly"]

    def month_sheet(title, value, fmt, rises=False):
        ws = wb.create_sheet(title)
        sheets[title] = ws
        header(ws, ["Year"] + months, {1: 8})
        for q in M:
            ws.append([q["year"]] + [value(q, c) for c in range(12)])
            for c in range(12):
                cell = ws.cell(row=ws.max_row, column=c + 2)
                cell.number_format = fmt
                if cell.value is None:
                    continue
                if rises and cell.value > 0:
                    cell.fill, cell.font = rise_fill, rise_font
                elif not rises and q["source"][c] == "interpolated":
                    cell.font = interp_font

    sheets = {}
    month_sheet("1st of month (m)", lambda q, c: q["level"][c], "0.000")
    month_sheet("Year-to-year (cm)", lambda q, c: q["yoy_cm"][c], "+0.0;-0.0;0.0", rises=True)
    month_sheet("Year-to-year (in)",
                lambda q, c: None if q["yoy_cm"][c] is None else round(q["yoy_cm"][c] / 2.54, 1), "+0.0;-0.0;0.0", rises=True)
    month_sheet("Total change (cm)", lambda q, c: q["total_cm"][c], "+0.0;-0.0;0.0")
    month_sheet("Total change (in)",
                lambda q, c: None if q["total_cm"][c] is None else round(q["total_cm"][c] / 2.54, 1), "+0.0;-0.0;0.0")

    # July 1 projection
    P = data["projection"]
    ws = wb.create_sheet("July 1 projection")
    ws.append([f"Projection for {P['target']}, starting from the {P['anchor']['date']} reading "
               f"({P['anchor']['level']:.3f} m), using {P['years_of_history']} years of history"])
    ws["A1"].font = bold
    ws.append([])
    ws.append(["Scenario", "Percentile", "Projected level (m)", "vs latest reading (cm)", "vs latest reading (in)",
               f"vs July 1, {P['last_july']['year']} (cm)", f"vs July 1, {P['last_july']['year']} (in)"])
    for c in range(1, 8):
        ws.cell(row=3, column=c).font, ws.cell(row=3, column=c).fill = bold, head_fill
        ws.column_dimensions[get_column_letter(c)].width = 18
    for sc in P["scenarios"]:
        ws.append([sc["name"], sc["percentile"] / 100, sc["level"], sc["change_cm"], round(sc["change_cm"] / 2.54, 1),
                   sc["vs_last_jul_cm"], round(sc["vs_last_jul_cm"] / 2.54, 1)])
        ws.cell(row=ws.max_row, column=2).number_format = "0%"
        ws.cell(row=ws.max_row, column=3).number_format = "0.000"
    ws.append([])
    ws.append(["July 1", "Trend, all years (m)", "Trend, last 10 years (m)", "Typical yearly change (m)"])
    for c in range(1, 5):
        ws.cell(row=ws.max_row, column=c).font, ws.cell(row=ws.max_row, column=c).fill = bold, head_fill
    for lt in P["long_term"]:
        ws.append([lt["year"], lt["trend_all"], lt["trend_10"], lt["typical"]])
        for c in (2, 3, 4):
            ws.cell(row=ws.max_row, column=c).number_format = "0.000"

    # data for the July 1 chart: actual July 1 levels, then the projections
    ws.append([])
    ws.append(["Year", "Actual July 1 (m)", "Trend, last 10 years (m)", "Typical yearly change (m)"])
    jstart = ws.max_row
    for c in range(1, 5):
        ws.cell(row=jstart, column=c).font, ws.cell(row=jstart, column=c).fill = bold, head_fill
    lastj = P["last_july"]
    for q in M:
        bridge = q["year"] == lastj["year"]  # join the projection lines to the last actual value
        ws.append([q["year"], q["level"][6], lastj["level"] if bridge else None, lastj["level"] if bridge else None])
    for lt in P["long_term"]:
        if lt["year"] > M[-1]["year"]:
            ws.append([lt["year"], None, lt["trend_10"], lt["typical"]])
    jend = ws.max_row
    for row in ws.iter_rows(min_row=jstart + 1, max_row=jend, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "0.000"
    ws_proj = ws

    # Charts
    cs = wb.create_sheet("Charts", 0)
    cs["A1"] = "Charts — Muriel Lake (06AC007). The tables they come from are on the other sheets."
    cs["A1"].font = Font(bold=True, size=13)
    cs.sheet_properties.pageSetUpPr.fitToPage = True      # prints/exports the charts on one page
    cs.page_setup.fitToWidth = cs.page_setup.fitToHeight = 1
    cs.page_setup.orientation = "landscape"
    BLUE, ORANGE, GREEN, PURPLE, GREY = "2F6FB3", "D9822B", "2E9A5B", "8B5CC4", "8A94A3"

    def tidy(ch, title, ytitle, xtitle="Year"):
        ch.title, ch.y_axis.title, ch.x_axis.title = title, ytitle, xtitle
        ch.width, ch.height = 24, 10
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.x_axis.tickLblPos = "low"          # keep year labels along the bottom, not on the zero line
        ch.y_axis.number_format = "0.0"
        if ch.legend is not None:
            ch.legend.position = "b"
        return ch

    def daily_chart(col, title, ytitle, colour):
        ch = tidy(LineChart(), title, ytitle, "Date")
        ch.y_axis.crossAx = 500
        ch.x_axis = DateAxis(crossAx=100)
        ch.x_axis.number_format, ch.x_axis.majorTimeUnit, ch.x_axis.title = "yyyy", "years", "Date"
        ch.x_axis.majorUnit, ch.x_axis.tickLblPos = 5, "low"
        ch.x_axis.delete = False
        ch.add_data(Reference(ws_daily, min_col=col, min_row=1, max_row=n_daily + 1), titles_from_data=True)
        ch.set_categories(Reference(ws_daily, min_col=1, min_row=2, max_row=n_daily + 1))
        ser = ch.series[0]
        ser.graphicalProperties.line.solidFill = colour
        ser.graphicalProperties.line.width = 15000
        ser.marker.symbol, ser.smooth = "none", False
        ch.legend = None
        return ch

    cs.add_chart(daily_chart(3, "Daily water level vs first reading (m)", "Change (m)", BLUE), "A3")
    cs.add_chart(daily_chart(4, "Daily water level vs first reading (ft)", "Change (ft)", ORANGE), "P3")

    ch = tidy(LineChart(), "July 1 water level: history and projections", "Level (m above sea level)")
    ch.add_data(Reference(ws_proj, min_col=2, max_col=4, min_row=jstart, max_row=jend), titles_from_data=True)
    ch.set_categories(Reference(ws_proj, min_col=1, min_row=jstart + 1, max_row=jend))
    for ser, colour, dash in zip(ch.series, (BLUE, PURPLE, GREY), (None, "dash", "dash")):
        ser.graphicalProperties.line.solidFill = colour
        ser.graphicalProperties.line.width = 22000
        if dash:
            ser.graphicalProperties.line.dashStyle = dash
        ser.marker.symbol, ser.smooth = "none", False
    ch.display_blanks = "gap"
    ch.x_axis.tickLblSkip = 5
    cs.add_chart(ch, "A24")

    tot = sheets["Total change (cm)"]
    ch = tidy(LineChart(), "Total change since first reading, 1st of Jan / Apr / Jul / Oct (cm)", "Change (cm)")
    for col, colour in zip((2, 5, 8, 11), (BLUE, ORANGE, GREEN, PURPLE)):
        ch.add_data(Reference(tot, min_col=col, min_row=1, max_row=tot.max_row), titles_from_data=True)
        ser = ch.series[-1]
        ser.graphicalProperties.line.solidFill = colour
        ser.graphicalProperties.line.width = 20000
        ser.marker.symbol, ser.smooth = "none", False
    ch.set_categories(Reference(tot, min_col=1, min_row=2, max_row=tot.max_row))
    ch.display_blanks = "gap"
    ch.x_axis.tickLblSkip = 5
    cs.add_chart(ch, "P24")

    yoy = sheets["Year-to-year (cm)"]
    for i, (col, name) in enumerate(zip((2, 5, 8, 11), ("Jan 1", "Apr 1", "Jul 1", "Oct 1"))):
        ch = tidy(BarChart(), f"{name}: change from the year before (cm) — green = rose", "Change (cm)")
        ch.type, ch.gapWidth = "col", 40
        ch.add_data(Reference(yoy, min_col=col, min_row=1, max_row=yoy.max_row), titles_from_data=True)
        ch.set_categories(Reference(yoy, min_col=1, min_row=2, max_row=yoy.max_row))
        ser = ch.series[0]
        ser.graphicalProperties.solidFill = ORANGE
        ser.invertIfNegative = False
        for idx, q in enumerate(M):
            v = q["yoy_cm"][col - 2]
            if v is not None and v > 0:
                pt = DataPoint(idx=idx)
                pt.graphicalProperties.solidFill = GREEN
                ser.dPt.append(pt)
        ch.legend = None
        ch.x_axis.tickLblSkip = 5
        cs.add_chart(ch, ("A", "P")[i % 2] + str(45 + 21 * (i // 2)))

    # About
    ws = wb.create_sheet("About")
    for line in [
        f"Muriel Lake water levels — Water Survey of Canada station {STATION} ({STATION_NAME})",
        f"Generated {data['updated']} (UTC). Latest reading {data['latest']['date']}: {data['latest']['level']:.3f} m.",
        f"Approved historical data (HYDAT) through {data['approved_until']}; later dates are provisional real-time data.",
        "Days with no reading are filled by straight-line interpolation (grey italic in the 1st-of-month table).",
        "Green cells in the year-to-year sheets mean the lake was higher than on the same date the year before.",
        "Data © Environment and Climate Change Canada, Open Government Licence – Canada.",
        "Live page: https://jaybeehan.github.io/muriel-lake/",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 110

    wb.save(XLSX)
    print(f"Wrote {XLSX} ({XLSX.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    build()
